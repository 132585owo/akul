import streamlit as st
import pypdf
import docx
from pptx import Presentation
import google.generativeai as genai

# 從 Streamlit 雲端安全保險箱讀取金鑰
genai.configure(api_key=st.secrets["GOOGLE_API_KEY"])
model = genai.GenerativeModel('gemini-3-flash')

# 定義萃取檔案文字的函式
def extract_text(file, file_type):
    text = ""
    try:
        if file_type == "pdf":
            pdf_reader = pypdf.PdfReader(file)
            for page in pdf_reader.pages:
                if page.extract_text():
                    text += page.extract_text()
        elif file_type == "docx":
            doc = docx.Document(file)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif file_type == "pptx":
            ppt = Presentation(file)
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
    except Exception as e:
        st.error(f"讀取檔案時發生錯誤: {e}")
    return text

# 網頁介面設計
st.set_page_config(page_title="AI 講義助教", page_icon="🤖")
st.title("🤖 專屬講義 AI 助教")
st.write("上傳每週講義並輸入作業，讓我直接幫你從講義中找答案！")

# 檔案與題目輸入區塊
uploaded_file = st.file_uploader("請上傳講義 (支援 PDF, Word, PPT)", type=["pdf", "docx", "pptx"])
question = st.text_area("請輸入本週作業題目：", height=150)

# 點擊執行按鈕後的處理
if st.button("開始解析與作答", type="primary"):
    if uploaded_file and question:
        with st.spinner("AI 正在努力閱讀講義中，請稍候..."):
            try:
                file_type = uploaded_file.name.split('.')[-1]
                content = extract_text(uploaded_file, file_type)
                
                prompt = f"""
                你現在是一位專業的學科助教。請仔細閱讀以下講義內容，並根據講義的脈絡與知識點，解答學生的作業題目。
                若有講義沒提到的部分，可以使用你的知識補充說明。
                
                【講義內容】：
                {content}
                
                【作業題目】：
                {question}
                """
                
                response = model.generate_content(prompt)
                
                st.success("解析完成！")
                st.write("### 📝 AI 助教的解答：")
                st.write(response.text)
                     
            except Exception as e:
                st.error(f"發生錯誤：{e}")
    else:
        st.warning("請記得上傳講義檔案，並輸入作業題目喔！")
